import json, df01
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
Q, YC = df01.Q, df01.YC
ids = ['JQ-%04d' % (25 + i) for i in range(48)]
gmap = {g: 'G-%04d' % (5 + k) for k, g in enumerate('ABCDEFGH')}
LISTEN = {'3-1', '3-2', '3-3'}
bank, rap, lms, media = [], [], [['Audio Url','Content','Correct Answer','Explanation','Image Url','Option 1','Option 2','Option 3','Option 4','STT','Tag']], []
gtext = {}
for i, (d, cap, bai, g, th, nd, ch, lc, da, gt, bh, sc) in enumerate(Q):
    vi, jp = YC[d]
    gid = gmap.get(g, '')
    if gid and nd and gid not in gtext: gtext[gid] = '\n'.join(l for l in nd.split('\n') if not l.startswith('['))
    img = (ids[i] if d in LISTEN or d == '1-1' else '')
    ve = (df01.VE + ('' if d == '1-1' else '') + bh) if bh else ''
    row = [ids[i], gid, 'DF-01', "'" + d, 'Đề full', cap, bai, '', 'DF-01 (Claude soạn)', jp, vi, '', th, nd, ch, len(lc)] + (lc + [''] * 4)[:4] + \
          [da, '', '', gt, bh, '', '', sc, '', 'Claude', 'Bổ sung theo chủ đề', 'Đạt', 'Claude soát (cần GV duyệt nhanh)', '', '', '', 'Chờ QC JP']
    assert len(row) == 37
    bank.append(row + ['', '', '', 'v1', '', '', ve])   # AL AM AN(skip) AO AP AQ AR
    rap.append(['DF-01', i + 1, ids[i], d, 'Đề full', 'Chờ QC JP'])
    lines = [vi] + ([jp] if jp else []) + ([th] if th else [])
    body = '\n'.join(l for l in nd.split('\n') if not l.startswith('[')) or (gtext.get(gid, '') if d in ('4-1', '4-2') else '')
    if body and d not in LISTEN: lines.append(body)
    if ch: lines.append(ch)
    aud = ('AUDIO_' + (gid or ids[i]) + '.wav') if d in LISTEN else ''
    lms.append([aud, '\n'.join(lines), da, gt, (ids[i] + '.png') if img else '', *(lc + [''] * 4)[:4], i + 1, 'JFT|%s|%s-%02d|%s' % (d, cap, bai, ids[i])])
    if img: media.append(['Tranh', ids[i] + '.png', d, ve])
    if d in LISTEN and sc: media.append(['Audio', 'AUDIO_' + (gid or ids[i]) + '.wav', d, sc])
json.dump({'bank': bank, 'rap': rap}, open('df01_rows.json', 'w'), ensure_ascii=False)
wb = Workbook(); ws = wb.active; ws.title = 'MULTIPLE CHOICE (Advanced)'
for r in lms: ws.append(r)
for c in ws[1]: c.font = Font(bold=True)
wb.save('../exports/DF-01_LMS_import.xlsx') if __import__('os').path.isdir('../exports') else None
import os; os.makedirs('../exports', exist_ok=True); wb.save('../exports/DF-01_LMS_import.xlsx')
wm = Workbook(); m = wm.active; m.title = 'Media DF-01'
m.append(['Loại', 'Tên file (đặt đúng)', 'Dạng', 'Prompt vẽ (dán ChatGPT) / Script audio (dán Gemini AI Studio)'])
for r in media: m.append(r)
for row in m.iter_rows(): [setattr(c, 'alignment', Alignment(wrap_text=True, vertical='top')) for c in row]
m.column_dimensions['B'].width = 22; m.column_dimensions['D'].width = 110
wm.save('../exports/DF-01_media_checklist.xlsx')
print(len(bank), len(media), sum(1 for x in media if x[0]=='Audio'), sum(1 for x in media if x[0]=='Tranh'))
