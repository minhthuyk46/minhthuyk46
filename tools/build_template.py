"""Sinh file mẫu Google Sheet cho hệ thống tạo đề JFT (phạm vi いろどり).
Chạy: python3 tools/build_template.py  → templates/JFT_Irodori_He_thong_tao_de.xlsx"""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter as L

import sys
LEAN = '--lean' in sys.argv  # bỏ công thức từng dòng (dùng ARRAYFORMULA trên Google Sheets)
OUT = 'templates/JFT_Irodori_He_thong_tao_de' + ('_lean' if LEAN else '') + '.xlsx'
NAVY, WHITE = '1F3A5F', 'FFFFFF'
FILL = {'common': 'DCE6F1', 'type': 'FDE9D9', 'auto': 'E2EFDA', 'ai': 'FFF2CC'}
thin = Side(style='thin', color='BFBFBF')

BOOKS = ['入門', '初級1', '初級2']
STATUS_REQ = ['Nháp', 'Sẵn sàng sinh', 'Đã sinh', 'Huỷ']
STATUS_Q = ['Draft', 'Đang duyệt VN', 'Đang duyệt JP', 'Cần sửa', 'Approved', 'Loại']
TOOLS = ['ChatGPT', 'Claude', 'GV tự soạn']
TYPES = [  # (mã sheet, phần thi, dạng bài)
    ('IN_01', '文字と語彙', '語の意味'), ('IN_02', '文字と語彙', '語の使い方'),
    ('IN_03', '文字と語彙', '漢字の読み'), ('IN_04', '文字と語彙', '漢字の表記'),
    ('IN_05', '会話と表現', '表現'), ('IN_06', '会話と表現', '文法'),
    ('IN_07', '聴解', '会話の聴解'), ('IN_08', '聴解', 'アナウンス'), ('IN_09', '聴解', '指示の聴解'),
    ('IN_10', '読解', '掲示の読解'), ('IN_11', '読解', '文章の読解'),
]

wb = Workbook()
wb.remove(wb.active)


def sheet(name, cols, rows=(), widths=None, groups=None, note=None, color='1F3A5F'):
    """cols: list[(header, ghi chú)] ; groups: list same length with key of FILL."""
    ws = wb.create_sheet(name)
    ws.sheet_properties.tabColor = color
    start = 1
    if note:
        ws.cell(1, 1, note).font = Font(italic=True, color='595959')
        start = 2
    for j, (h, tip) in enumerate(cols, 1):
        c = ws.cell(start, j, h)
        g = groups[j - 1] if groups else None
        c.font = Font(bold=True, color=WHITE if not g else '000000')
        c.fill = PatternFill('solid', fgColor=NAVY if not g else FILL[g])
        c.alignment = Alignment(wrap_text=True, vertical='center', horizontal='center')
        c.border = Border(top=thin, bottom=thin, left=thin, right=thin)
        if tip and not LEAN:
            from openpyxl.comments import Comment
            c.comment = Comment(tip, 'JFT')
        ws.column_dimensions[L(j)].width = (widths or {}).get(h, max(12, min(40, len(h) * 2 + 4)))
    for i, r in enumerate(rows, start + 1):
        for j, v in enumerate(r, 1):
            ws.cell(i, j, v).alignment = Alignment(wrap_text=True, vertical='top')
    ws.row_dimensions[start].height = 42
    ws.freeze_panes = ws.cell(start + 1, 3)
    return ws, start


def dropdown(ws, col, values_or_ref, first_row, last_row=500):
    f = values_or_ref if values_or_ref.startswith('=') else '"' + ','.join(values_or_ref.split('|')) + '"'
    dv = DataValidation(type='list', formula1=f, allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f'{L(col)}{first_row}:{L(col)}{last_row}')


# ---------------- 00 HƯỚNG DẪN ----------------
ws = wb.create_sheet('00_HUONG_DAN'); ws.sheet_properties.tabColor = '7F7F7F'
guide = [
    ('HỆ THỐNG TẠO ĐỀ JFT-BASIC — PHẠM VI いろどり', ''),
    ('', ''),
    ('FLOW', 'Kế hoạch → Nhập nguồn → Nhập yêu cầu theo dạng bài → AI sinh → Duyệt → Ráp đề → Xuất LMS → Thống kê'),
    ('', ''),
    ('Nhóm sheet', 'Vai trò'),
    ('01_KE_HOACH', 'Lên kế hoạch từng đề thi thử: ngày thi, phạm vi bài, số câu mỗi dạng, người phụ trách'),
    ('02_THONG_KE', 'Tự động: số câu theo dạng × trạng thái, độ phủ theo bài, đủ/thiếu so với kế hoạch'),
    ('10–13_NGUON_*', 'Dữ liệu nguồn từ giáo trình Irodori: từ vựng, kanji, ngữ pháp, tình huống. AI chỉ được dùng trong phạm vi này'),
    ('IN_01 → IN_11', 'Phiếu yêu cầu sinh câu cho từng dạng bài. 1 dòng = 1 yêu cầu (có thể sinh nhiều câu)'),
    ('20_PROMPT', 'Mẫu prompt cho từng dạng bài, Apps Script thay {biến} bằng dữ liệu dòng IN_xx'),
    ('30_NGAN_HANG', 'Ngân hàng câu hỏi sau khi AI sinh + cột duyệt VN/JP'),
    ('40_RAP_DE', 'Danh sách câu đã chọn cho từng mã đề → xuất file import LMS'),
    ('99_DANH_MUC', 'Danh mục dùng cho dropdown. Sửa ở đây, không sửa trong ô'),
    ('', ''),
    ('Màu tiêu đề cột', ''),
    ('Xanh dương nhạt', 'Trường chung — bắt buộc ở mọi dạng bài'),
    ('Cam nhạt', 'Trường riêng của dạng bài'),
    ('Vàng', 'Gợi ý cho AI (không bắt buộc, để trống thì AI tự chọn)'),
    ('Xanh lá', 'Tự động — KHÔNG nhập tay (công thức / Apps Script ghi)'),
    ('', ''),
    ('Quy tắc nhập', ''),
    ('1', 'Tiếng Nhật nhập đúng dạng sẽ xuất hiện trong đề (kanji + furigana theo chuẩn JFT nếu cần)'),
    ('2', 'Nhiều giá trị trong 1 ô: ngăn cách bằng dấu ;  (ví dụ: 朝ご飯;昼ご飯;晩ご飯)'),
    ('3', 'Chỉ đổi Trạng thái = "Sẵn sàng sinh" khi đã điền đủ trường bắt buộc'),
    ('4', 'Mỗi câu AI sinh luôn vào Draft — chỉ câu Approved mới được ráp đề'),
    ('5', 'Dòng có chữ VÍ DỤ ở cột Ghi chú là dòng mẫu minh hoạ — xoá trước khi chạy thật'),
]
for i, (a, b) in enumerate(guide, 1):
    ws.cell(i, 1, a); ws.cell(i, 2, b)
    if a in ('FLOW', 'Nhóm sheet', 'Màu tiêu đề cột', 'Quy tắc nhập') or i == 1:
        ws.cell(i, 1).font = Font(bold=True, size=13 if i == 1 else 11, color=NAVY)
for i, k in zip(range(16, 20), ['common', 'type', 'ai', 'auto']):
    ws.cell(i, 1).fill = PatternFill('solid', fgColor=FILL[k])
ws.column_dimensions['A'].width = 22; ws.column_dimensions['B'].width = 110

# ---------------- 99 DANH MỤC ----------------
dm, _ = sheet('99_DANH_MUC', [('Quyển', ''), ('Phần thi', ''), ('Dạng bài', ''), ('Mã sheet input', ''),
                              ('TT yêu cầu', ''), ('TT câu hỏi', ''), ('Công cụ sinh', ''),
                              ('Kiểu đáp án', ''), ('Mức lịch sự', ''), ('Giọng', '')], color='7F7F7F')
lists = [BOOKS, [t[1] for t in TYPES], [t[2] for t in TYPES], [t[0] for t in TYPES], STATUS_REQ, STATUS_Q,
         TOOLS, ['Chữ', 'Hình 4 ô'], ['普通形', 'です・ます', '敬語(尊敬・謙譲)'], ['F', 'M', 'N (dẫn)']]
for j, lst in enumerate(lists, 1):
    for i, v in enumerate(lst, 2):
        dm.cell(i, j, v)
REF = lambda col, n: f"='99_DANH_MUC'!${col}$2:${col}${n + 1}"

# ---------------- 01 KẾ HOẠCH ----------------
plan_cols = [('Mã đề', 'VD: TT-2026-01'), ('Tên đề', ''), ('Lớp / đối tượng', ''), ('Ngày thi', ''),
             ('Quyển', ''), ('Bài từ', ''), ('Bài đến', '')] + \
            [(f'{t[2]}', f'Số câu {t[2]} trong đề') for t in TYPES] + \
            [('Tổng câu', 'Tự tính'), ('Phụ trách', ''), ('Deadline duyệt', ''), ('Trạng thái', ''), ('Ghi chú', '')]
g = ['common'] * 7 + ['type'] * len(TYPES) + ['auto', 'common', 'common', 'common', 'common']
ws, h = sheet('01_KE_HOACH', plan_cols, groups=g, color='2E75B6',
              note='Mỗi dòng = 1 đề thi thử. Số câu mỗi dạng lấy theo cấu trúc đề JFT chính thức (đối chiếu đề mẫu Japan Foundation).')
ws.cell(h + 1, 1, 'TT-2026-01'); ws.cell(h + 1, 2, 'Thi thử giữa khoá (VÍ DỤ)'); ws.cell(h + 1, 5, '入門')
ws.cell(h + 1, 6, 1); ws.cell(h + 1, 7, 9)
for r in ([] if LEAN else range(h + 1, 201)):
    ws.cell(r, 8 + len(TYPES), f'=IF(A{r}="","",SUM({L(8)}{r}:{L(7 + len(TYPES))}{r}))')
dropdown(ws, 5, REF('A', 3), h + 1, 200)
dropdown(ws, 8 + len(TYPES) + 3, 'Lên kế hoạch|Đang soạn|Đang duyệt|Đã xuất LMS|Đã thi', h + 1, 200)

# ---------------- NGUỒN ----------------
src_common = [('Quyển', ''), ('Bài', 'Số bài (第○課)'), ('Chủ đề bài', '')]
sheet('10_NGUON_TU_VUNG', src_common + [
    ('Từ (dạng xuất hiện trong sách)', 'Kanji/kana đúng như giáo trình'), ('Cách đọc (hiragana)', ''),
    ('Nghĩa tiếng Việt', ''), ('Từ loại', 'Danh/Động/Tính い/Tính な/Phó…'), ('Nhóm nghĩa', 'Đồ ăn, giao thông, công việc…'),
    ('Vẽ hình được?', 'Y/N — dùng cho 語の意味 dạng hình'), ('Từ dễ nhầm', 'Từ gần nghĩa/gần âm, ngăn cách ;'),
    ('Đã dùng (lần)', 'Tự động')], groups=['common'] * 3 + ['type'] * 7 + ['auto'], color='548235')
sheet('11_NGUON_KANJI', src_common + [
    ('Kanji', ''), ('Từ chứa kanji (trong sách)', 'ngăn cách ;'), ('Cách đọc đúng của từ', 'ngăn cách ; theo thứ tự cột trước'),
    ('Âm On', ''), ('Âm Kun', ''), ('Kanji giống hình', 'VD: 休↔体, 未↔末'), ('Đã dùng (lần)', 'Tự động')],
    groups=['common'] * 3 + ['type'] * 6 + ['auto'], color='548235')
sheet('12_NGUON_NGU_PHAP', src_common + [
    ('Mẫu ngữ pháp', 'VD: 〜てください'), ('Chức năng', 'Nhờ vả, xin phép, mời…'), ('Câu ví dụ trong sách', ''),
    ('Trợ từ / mẫu dễ nhầm', 'Dùng làm phương án nhiễu, ngăn cách ;'), ('Đã dùng (lần)', 'Tự động')],
    groups=['common'] * 3 + ['type'] * 4 + ['auto'], color='548235')
sheet('13_NGUON_TINH_HUONG', src_common + [
    ('Mã tình huống', 'VD: TH-入門-05-01'), ('Can-do (いろどり)', 'Chép đúng can-do của bài'), ('Địa điểm', 'Cửa hàng, ga, công ty…'),
    ('Nhân vật / quan hệ', 'VD: nhân viên mới – sếp'), ('Từ khoá / mẫu câu chính', ''),
    ('Dùng cho dạng', '表現/聴解/読解, ngăn cách ;'), ('Đã dùng (lần)', 'Tự động')],
    groups=['common'] * 3 + ['type'] * 6 + ['auto'], color='548235')
for n in ('10_NGUON_TU_VUNG', '11_NGUON_KANJI', '12_NGUON_NGU_PHAP', '13_NGUON_TINH_HUONG'):
    dropdown(wb[n], 1, REF('A', 3), 2)

# ---------------- IN_xx: PHIẾU YÊU CẦU THEO DẠNG BÀI ----------------
COMMON = [('Mã yêu cầu', 'Tự sinh: IN_xx-0001'), ('Mã đề (kế hoạch)', 'Để trống nếu sinh vào ngân hàng chung'),
          ('Quyển', ''), ('Bài', ''), ('Số câu cần sinh', '1–5 / lần chạy')]
TAIL = [('Ghi chú cho AI', 'Yêu cầu thêm, VD: tránh chủ đề đồ ăn'), ('Công cụ sinh', ''), ('Người nhập', ''),
        ('Trạng thái', ''), ('Số câu đã sinh', 'Tự động'), ('Số câu Approved', 'Tự động')]

SPEC = {
 '語の意味': ([('Từ đích', 'Lấy từ 10_NGUON_TU_VUNG'), ('Nghĩa tiếng Việt', '')],
            [('Kiểu đáp án', 'Chữ / Hình 4 ô'), ('Ngữ cảnh câu hỏi', 'Câu chứa ___ hoặc mô tả hình'), ('Từ nhiễu gợi ý', 'ngăn cách ;')],
            ['語の意味', '会社', 'công ty', 'Hình 4 ô', 'わたしの ___ は 大きいです。', '学校;病院;銀行']),
 '語の使い方': ([('Từ đích', ''), ('Câu ngữ cảnh (có ___)', 'Câu chứa chỗ trống')],
             [('Collocation đúng', 'VD: 電話を かける'), ('Lỗi dùng từ muốn bẫy', 'VD: nhầm 着る/はく')],
             ['語の使い方', 'かける', '友だちに 電話を ___。', '電話を かける', 'する;とる;いう']),
 '漢字の読み': ([('Từ chứa kanji (gạch chân)', ''), ('Cách đọc đúng', 'hiragana')],
             [('Câu ngữ cảnh', 'Câu chứa từ đích'), ('Kiểu nhiễu', 'Trường âm / âm đục / âm ngắt / On-Kun, ngăn cách ;')],
             ['漢字の読み', '会社', 'かいしゃ', 'あした 会社に 行きます。', 'Trường âm;Âm đục']),
 '漢字の表記': ([('Từ viết hiragana (gạch chân)', ''), ('Kanji đúng', '')],
             [('Câu ngữ cảnh', ''), ('Kanji nhiễu', 'Giống hình / cùng âm, ngăn cách ;')],
             ['漢字の表記', 'やすみ', '休み', 'あしたは やすみです。', '体み;休見;安み']),
 '表現': ([('Mã tình huống', 'Lấy từ 13_NGUON_TINH_HUONG'), ('Chức năng giao tiếp', 'Xin phép, cảm ơn, xin lỗi, mời…'),
          ('Người nói → người nghe', 'VD: nhân viên → sếp'), ('Mức lịch sự', '')],
         [('Mô tả tình huống (VN hoặc JP)', ''), ('Câu đúng mong muốn', ''), ('Câu nhiễu gợi ý', 'Sai chức năng / sai mức lịch sự')],
         ['表現', 'TH-VÍ DỤ', 'Xin phép về trước', 'nhân viên → sếp', 'です・ます', 'Hết giờ làm, muốn về trước sếp',
          'お先に失礼します。', 'お疲れさまでした。;いってきます。']),
 '文法': ([('Mẫu ngữ pháp đích', 'Lấy từ 12_NGUON_NGU_PHAP'), ('Dạng câu', 'Hội thoại A-B / câu đơn')],
         [('Ngữ cảnh', ''), ('Vị trí ô trống', 'Trợ từ / đuôi động từ / liên từ'), ('Mẫu dễ nhầm làm nhiễu', 'ngăn cách ;')],
         ['文法', '〜てください', 'Hội thoại A-B', 'Ở văn phòng, sếp nhờ nhân viên', 'Đuôi động từ', '〜てもいいです;〜ています']),
 '会話の聴解': ([('Mã tình huống', ''), ('Địa điểm', ''), ('Người nói', 'VD: F=nhân viên cửa hàng; M=khách'),
              ('Câu hỏi trọng tâm', 'Hỏi cái gì? Mấy giờ? Ở đâu?'), ('Kiểu đáp án', 'Chữ / Hình 4 ô')],
             [('Thông tin bẫy', 'Thông tin được nhắc rồi bị đổi/phủ định'), ('Số lượt thoại', '4–8'), ('Từ vựng/ngữ pháp bắt buộc dùng', '')],
             ['会話の聴解', 'TH-VÍ DỤ', 'Cửa hàng tiện lợi', 'F=nhân viên; M=khách', 'Khách mua gì?', 'Hình 4 ô',
              'Định mua cơm hộp nhưng đổi sang bánh mì', 6, 'おにぎり;パン']),
 'アナウンス': ([('Địa điểm', 'Ga, siêu thị, công ty, khu dân cư'), ('Loại thông báo', 'Trễ tàu, giảm giá, sự cố, lịch họp'),
              ('Thông tin cần nghe', 'Giờ / số / địa điểm / hành động')],
             [('Giọng', 'F/M'), ('Thông tin bẫy', ''), ('Độ dài (giây)', '15–40')],
             ['アナウンス', 'Ga tàu', 'Tàu trễ', 'Tàu đến lúc mấy giờ', 'F', 'Giờ cũ được nhắc trước giờ mới', 25]),
 '指示の聴解': ([('Người chỉ thị → người nhận', ''), ('Nhiệm vụ', 'VD: chuẩn bị phòng họp'), ('Câu hỏi', 'Phải làm gì trước tiên?')],
             [('Số bước trong chỉ thị', '2–4'), ('Thông tin bẫy', 'Bước bị huỷ / đổi thứ tự'), ('Kiểu đáp án', 'Chữ / Hình 4 ô')],
             ['指示の聴解', 'sếp → nhân viên mới', 'Chuẩn bị phòng họp', 'Đầu tiên phải làm gì?', 3, 'Bước photo tài liệu bị huỷ', 'Hình 4 ô']),
 '掲示の読解': ([('Loại văn bản', 'Biển báo, poster, lịch, tin nhắn, thực đơn, nội quy'), ('Bối cảnh', ''), ('Câu hỏi', '')],
             [('Thông tin chính cần có', 'Ngày, giờ, giá, điều kiện… ngăn cách ;'), ('Thông tin nhiễu', ''), ('Cần hình?', 'Y/N')],
             ['掲示の読解', 'Thông báo đổ rác', 'Khu chung cư', 'Rác cháy được bỏ ngày nào?', 'Thứ 2;Thứ 5;trước 8 giờ', 'Ngày rác nhựa', 'Y']),
 '文章の読解': ([('Thể loại', 'Email, blog, thông báo, bài viết'), ('Chủ đề', ''), ('Số câu hỏi con', '1–3, mỗi câu thành 1 dòng LMS')],
             [('Độ dài (字)', '150–400'), ('Người viết → người đọc', ''), ('Ý chính cần hỏi', 'ngăn cách ;')],
             ['文章の読解', 'Email', 'Mời tiệc chia tay đồng nghiệp', 2, 250, 'đồng nghiệp → cả phòng', 'Thời gian;Cần mang gì']),
}

for code, sec, typ in TYPES:
    req, opt, sample = SPEC[typ]
    cols = COMMON + req + opt + TAIL
    groups = ['common'] * len(COMMON) + ['type'] * len(req) + ['ai'] * len(opt) + \
             ['common', 'common', 'common', 'common', 'auto', 'auto']
    ws, h = sheet(f'{code}_{typ}', cols, groups=groups, color='C55A11',
                  note=f'{sec} › {typ}  |  1 dòng = 1 yêu cầu sinh câu. Cam = bắt buộc, vàng = gợi ý cho AI.')
    r = h + 1
    ws.cell(r, 1, f'{code}-0001'); ws.cell(r, 3, '入門'); ws.cell(r, 4, 1); ws.cell(r, 5, 3)
    for j, v in enumerate(sample[1:], len(COMMON) + 1):
        ws.cell(r, j, v)
    n = len(cols)
    ws.cell(r, n - 5, 'VÍ DỤ — xoá dòng này khi dùng thật'); ws.cell(r, n - 4, 'ChatGPT'); ws.cell(r, n - 2, 'Nháp')
    for rr in ([] if LEAN else range(h + 1, 301)):
        ws.cell(rr, n - 1, f"=IF(A{rr}=\"\",\"\",COUNTIF('30_NGAN_HANG'!$B:$B,A{rr}))")
        ws.cell(rr, n, f"=IF(A{rr}=\"\",\"\",COUNTIFS('30_NGAN_HANG'!$B:$B,A{rr},'30_NGAN_HANG'!$Y:$Y,\"Approved\"))")
    dropdown(ws, 3, REF('A', 3), h + 1, 300)
    dropdown(ws, n - 4, REF('G', 3), h + 1, 300)
    dropdown(ws, n - 2, REF('E', 4), h + 1, 300)
    for j, (hd, _) in enumerate(cols, 1):
        if hd == 'Kiểu đáp án':
            dropdown(ws, j, REF('H', 2), h + 1, 300)
        if hd == 'Mức lịch sự':
            dropdown(ws, j, REF('I', 3), h + 1, 300)

# ---------------- 20 PROMPT ----------------
P_COMMON = ('あなたはJFT-Basicの作問者です。出題範囲は『いろどり』{Quyển} 第{Bài}課まで。'
            '範囲外の語彙・漢字・文法は使わない。4択・正答1つ。正答位置はばらつかせる。'
            '解説はベトナム語で、正答の理由と各誤答の理由を書く。出力はJSON配列のみ。')
P = {
 '語の意味': '「{Từ đích}」（{Nghĩa tiếng Việt}）の意味を問う問題を{Số câu cần sinh}問。形式: {Kiểu đáp án}。文脈: {Ngữ cảnh câu hỏi}。誤答候補: {Từ nhiễu gợi ý}',
 '語の使い方': '「{Từ đích}」を正しく使えるか問う問題。文: {Câu ngữ cảnh (có ___)}。正しいコロケーション: {Collocation đúng}。狙う誤用: {Lỗi dùng từ muốn bẫy}',
 '漢字の読み': '下線部「{Từ chứa kanji (gạch chân)}」の読み（正答: {Cách đọc đúng}）を問う。文: {Câu ngữ cảnh}。誤答の作り方: {Kiểu nhiễu}',
 '漢字の表記': '下線部「{Từ viết hiragana (gạch chân)}」の漢字（正答: {Kanji đúng}）を問う。文: {Câu ngữ cảnh}。誤答候補: {Kanji nhiễu}',
 '表現': '場面: {Mô tả tình huống (VN hoặc JP)}。{Người nói → người nghe}、機能: {Chức năng giao tiếp}、丁寧さ: {Mức lịch sự}。適切な発話を選ぶ問題。正答例: {Câu đúng mong muốn}。誤答候補: {Câu nhiễu gợi ý}',
 '文法': '文法項目「{Mẫu ngữ pháp đích}」。形式: {Dạng câu}。場面: {Ngữ cảnh}。空欄位置: {Vị trí ô trống}。誤答に使う項目: {Mẫu dễ nhầm làm nhiễu}',
 '会話の聴解': '聴解スクリプトを書く（話者ラベル F:/M:/N:）。場所: {Địa điểm}。話者: {Người nói}。質問: {Câu hỏi trọng tâm}。形式: {Kiểu đáp án}。ひっかけ: {Thông tin bẫy}。{Số lượt thoại}ターン程度。必ず使う語: {Từ vựng/ngữ pháp bắt buộc dùng}',
 'アナウンス': '{Địa điểm}での{Loại thông báo}のアナウンス（N:）を{Độ dài (giây)}秒程度で書く。聞き取らせる情報: {Thông tin cần nghe}。ひっかけ: {Thông tin bẫy}',
 '指示の聴解': '{Người chỉ thị → người nhận}への指示の会話。タスク: {Nhiệm vụ}。手順{Số bước trong chỉ thị}つ。質問: {Câu hỏi}。ひっかけ: {Thông tin bẫy}。形式: {Kiểu đáp án}',
 '掲示の読解': '{Bối cảnh}の{Loại văn bản}を作成し、読解問題を作る。必要情報: {Thông tin chính cần có}。紛らわしい情報: {Thông tin nhiễu}。質問: {Câu hỏi}。画像化する場合はImagePromptに英語で指示',
 '文章の読解': '{Người viết → người đọc}の{Thể loại}（約{Độ dài (字)}字、テーマ: {Chủ đề}）を書き、{Số câu hỏi con}問作る。問う内容: {Ý chính cần hỏi}。各問は同じ本文を共有する',
}
ws, h = sheet('20_PROMPT', [('Dạng bài', ''), ('Prompt chung (system)', ''), ('Prompt riêng (user) — {biến} = tên cột IN_xx', ''),
                            ('Định dạng JSON đầu ra', ''), ('Phiên bản', ''), ('Ghi chú cải tiến', '')], color='BF8F00',
              widths={'Prompt chung (system)': 60, 'Prompt riêng (user) — {biến} = tên cột IN_xx': 80, 'Định dạng JSON đầu ra': 60})
JSON_FMT = '[{"Stem":"","Script":"","Option1":"","Option2":"","Option3":"","Option4":"","Answer":1,"Explanation_VI":"","ImagePrompt":""}]'
for i, (_, _, typ) in enumerate(TYPES, h + 1):
    for j, v in enumerate([typ, P_COMMON, P[typ], JSON_FMT, 'v1', ''], 1):
        ws.cell(i, j, v).alignment = Alignment(wrap_text=True, vertical='top')

# ---------------- 30 NGÂN HÀNG ----------------
bank_cols = [('Mã câu', 'Tự sinh'), ('Mã yêu cầu', 'Liên kết IN_xx'), ('Phần thi', ''), ('Dạng bài', ''), ('Quyển', ''), ('Bài', ''),
             ('Nội dung câu hỏi (Content)', '→ cột Content LMS'), ('Đoạn văn / Script nghe', 'Script: F:/M:/N: mỗi lượt 1 dòng'),
             ('Option 1', ''), ('Option 2', ''), ('Option 3', ''), ('Option 4', ''), ('Đáp án (1–4)', '→ Correct Answer LMS'),
             ('Giải thích (VI)', '→ Explanation LMS'), ('Image prompt', ''), ('Image Url', ''), ('Audio Url', ''),
             ('Công cụ sinh', ''), ('Ngày sinh', ''), ('Người duyệt VN', ''), ('Nhận xét VN', ''), ('Người duyệt JP', ''),
             ('Nhận xét JP', ''), ('Loại lỗi', 'Ngoài phạm vi / 2 đáp án đúng / sai tiếng Nhật / nhiễu yếu / lộ đáp án'),
             ('Trạng thái', ''), ('Số lần đã ra đề', 'Tự động'), ('Tag LMS', 'Tự động')]
g = ['auto', 'auto'] + ['common'] * 4 + ['ai'] * 11 + ['auto', 'auto'] + ['common'] * 6 + ['auto', 'auto']
ws, h = sheet('30_NGAN_HANG', bank_cols, groups=g, color='7030A0',
              widths={'Nội dung câu hỏi (Content)': 40, 'Đoạn văn / Script nghe': 45, 'Giải thích (VI)': 45})
dropdown(ws, 25, REF('F', 6), h + 1, 5000)
dropdown(ws, 24, 'Ngoài phạm vi|2 đáp án đúng|Sai tiếng Nhật|Nhiễu yếu|Lộ đáp án|Hình/Audio lỗi|Khác', h + 1, 5000)
dropdown(ws, 13, '1|2|3|4', h + 1, 5000)
for r in ([] if LEAN else range(h + 1, 2001)):
    ws.cell(r, 26, f"=IF(A{r}=\"\",\"\",COUNTIF('40_RAP_DE'!$C:$C,A{r}))")
    ws.cell(r, 27, f'=IF(A{r}="","","JFT|"&C{r}&"|"&D{r}&"|"&E{r}&"L"&F{r})')

# ---------------- 40 RÁP ĐỀ ----------------
ws, h = sheet('40_RAP_DE', [('Mã đề', ''), ('STT', ''), ('Mã câu', ''), ('Dạng bài', 'Tự động'),
                            ('Trạng thái câu', 'Tự động — phải là Approved'), ('Ghi chú', '')],
              groups=['common', 'common', 'common', 'auto', 'auto', 'common'], color='7030A0')
for r in ([] if LEAN else range(h + 1, 1001)):
    ws.cell(r, 4, f"=IF(C{r}=\"\",\"\",IFERROR(INDEX('30_NGAN_HANG'!$D:$D,MATCH(C{r},'30_NGAN_HANG'!$A:$A,0)),\"Không tìm thấy\"))")
    ws.cell(r, 5, f"=IF(C{r}=\"\",\"\",IFERROR(INDEX('30_NGAN_HANG'!$Y:$Y,MATCH(C{r},'30_NGAN_HANG'!$A:$A,0)),\"\"))")

# ---------------- 02 THỐNG KÊ ----------------
ws = wb.create_sheet('02_THONG_KE', 1); ws.sheet_properties.tabColor = '2E75B6'
ws['A1'] = 'A. Ngân hàng câu hỏi theo dạng bài × trạng thái'; ws['A1'].font = Font(bold=True, color=NAVY, size=12)
heads = ['Phần thi', 'Dạng bài'] + STATUS_Q + ['Tổng', 'Yêu cầu chưa sinh']
for j, t in enumerate(heads, 1):
    c = ws.cell(2, j, t); c.font = Font(bold=True, color=WHITE); c.fill = PatternFill('solid', fgColor=NAVY)
for i, (code, sec, typ) in enumerate(TYPES, 3):
    ws.cell(i, 1, sec); ws.cell(i, 2, typ)
    for j, st in enumerate(STATUS_Q, 3):
        ws.cell(i, j, f"=COUNTIFS('30_NGAN_HANG'!$D:$D,$B{i},'30_NGAN_HANG'!$Y:$Y,\"{st}\")")
    ws.cell(i, 3 + len(STATUS_Q), f"=COUNTIF('30_NGAN_HANG'!$D:$D,$B{i})")
    ws.cell(i, 4 + len(STATUS_Q), f"=COUNTIF('{code}_{typ}'!$A:$Z,\"Sẵn sàng sinh\")")
end = 2 + len(TYPES)
r0 = end + 3
ws.cell(r0, 1, 'B. Đủ / thiếu câu Approved cho 1 mã đề').font = Font(bold=True, color=NAVY, size=12)
ws.cell(r0 + 1, 1, 'Nhập mã đề →'); ws.cell(r0 + 1, 2, 'TT-2026-01'); ws.cell(r0 + 1, 2).fill = PatternFill('solid', fgColor=FILL['ai'])
for j, t in enumerate(['Dạng bài', 'Cần (kế hoạch)', 'Approved hiện có', 'Thiếu'], 1):
    c = ws.cell(r0 + 2, j, t); c.font = Font(bold=True, color=WHITE); c.fill = PatternFill('solid', fgColor=NAVY)
for k, (_, _, typ) in enumerate(TYPES):
    r = r0 + 3 + k
    ws.cell(r, 1, typ)
    ws.cell(r, 2, f"=IFERROR(INDEX('01_KE_HOACH'!{L(8 + k)}:{L(8 + k)},MATCH($B${r0 + 1},'01_KE_HOACH'!$A:$A,0)),0)")
    ws.cell(r, 3, f"=COUNTIFS('30_NGAN_HANG'!$D:$D,A{r},'30_NGAN_HANG'!$Y:$Y,\"Approved\")")
    ws.cell(r, 4, f'=MAX(0,B{r}-C{r})')
r1 = r0 + 4 + len(TYPES) + 1
ws.cell(r1, 1, 'C. Độ phủ theo bài (số câu Approved)').font = Font(bold=True, color=NAVY, size=12)
hd = ['Quyển', 'Bài'] + [t[2] for t in TYPES] + ['Tổng']
for j, t in enumerate(hd, 1):
    c = ws.cell(r1 + 1, j, t); c.font = Font(bold=True, color=WHITE); c.fill = PatternFill('solid', fgColor=NAVY)
r = r1 + 2
for b in ([] if LEAN else BOOKS):
    for les in range(1, 19):
        ws.cell(r, 1, b); ws.cell(r, 2, les)
        for k, (_, _, typ) in enumerate(TYPES, 3):
            ws.cell(r, k, f"=COUNTIFS('30_NGAN_HANG'!$E:$E,$A{r},'30_NGAN_HANG'!$F:$F,$B{r},'30_NGAN_HANG'!$D:$D,\"{typ}\",'30_NGAN_HANG'!$Y:$Y,\"Approved\")")
        ws.cell(r, 3 + len(TYPES), f'=SUM(C{r}:{L(2 + len(TYPES))}{r})')
        r += 1
ws.cell(r1 + 1, 1).comment = None
ws.column_dimensions['A'].width = 16; ws.column_dimensions['B'].width = 16
for j in range(3, 16):
    ws.column_dimensions[L(j)].width = 13
ws.freeze_panes = 'C3'

# Thứ tự tab
order = ['00_HUONG_DAN', '01_KE_HOACH', '02_THONG_KE', '10_NGUON_TU_VUNG', '11_NGUON_KANJI', '12_NGUON_NGU_PHAP',
         '13_NGUON_TINH_HUONG'] + [f'{c}_{t}' for c, _, t in TYPES] + ['20_PROMPT', '30_NGAN_HANG', '40_RAP_DE', '99_DANH_MUC']
wb._sheets = [wb[n] for n in order]
wb.save(OUT)
print('Saved', OUT)
