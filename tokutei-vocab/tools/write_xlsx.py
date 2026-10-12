"""python3 tools/write_xlsx.py out.json — ghi file .xlsx theo 2 template LMS."""
import json, sys, copy
from pathlib import Path
import openpyxl
from openpyxl.styles import Alignment
R = Path(__file__).resolve().parent.parent
data = json.load(open(sys.argv[1], encoding='utf-8'))
for les in data:
    slug = les['lesson'].replace('Bài', 'Bai').replace(' ', '_')
    wb = openpyxl.load_workbook(R / 'templates/Flashcard_tu_vung_bai_xx.xlsx'); ws = wb.active
    ws.delete_rows(2, ws.max_row)
    for r in les['flash']:
        ws.append(r)
        ws.cell(ws.max_row, 4).alignment = Alignment(wrap_text=True, vertical='center')
    wb.save(R / f'exports/flashcard/Flashcard_tu_vung_{slug}.xlsx')
    wb = openpyxl.load_workbook(R / 'templates/Form_lms-Quiz_bai_xx.xlsx'); ws = wb['MULTIPLE CHOICE (Advanced)']
    ws.delete_rows(2, ws.max_row)
    for r in les['quiz']:
        ws.append([v if v != '' else None for v in r])
    wb.save(R / f'exports/quiz/Quiz_tu_vung_{slug}.xlsx')
    print(les['lesson'], len(les['flash']), 'flashcard |', len(les['quiz']), 'câu quiz')
